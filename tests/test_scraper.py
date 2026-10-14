import sys
from argparse import Namespace
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from bot import scraper
from bot.scraper import (
    PortalBlockedError,
    detect_block_reason,
    extract_afp,
    format_rut_for_portal,
    get_public_ip,
    normalize_rut,
)
from desktop_backend import worker


def test_normalize_rut_con_puntos_y_guion():
    assert normalize_rut("15.800.185-3") == "158001853"


def test_normalize_rut_solo_guion():
    assert normalize_rut("15800185-3") == "158001853"


def test_normalize_rut_ya_limpio():
    assert normalize_rut("158001853") == "158001853"


def test_normalize_rut_strips_espacios():
    assert normalize_rut("  15800185-3  ") == "158001853"


def test_format_rut_for_portal():
    assert format_rut_for_portal("15.800.185-3") == "15800185-3"


def test_extract_afp_habitat():
    texto = (
        "Certifico que el(la) señor(a) RODRIGO IGNACIO ESCANILLA MIRANDA, "
        "RUT N° 15800185-3 se encuentra incorporado(a) a AFP HABITAT, "
        "con fecha 1 de Marzo de 2006."
    )
    assert extract_afp(texto) == "HABITAT"


def test_extract_afp_capital():
    texto = "...se encuentra incorporado(a) a AFP CAPITAL, con fecha..."
    assert extract_afp(texto) == "CAPITAL"


def test_extract_afp_planvital():
    texto = "...se encuentra incorporado(a) a AFP PLANVITAL, con fecha..."
    assert extract_afp(texto) == "PLANVITAL"


def test_extract_afp_no_encontrado():
    texto = "RUT no registrado o no encontrado en el sistema."
    assert extract_afp(texto) is None


def test_get_public_ip_retorna_string():
    mock_resp = MagicMock()
    mock_resp.read.return_value = b"1.2.3.4"
    mock_resp.__enter__ = lambda s: s
    mock_resp.__exit__ = MagicMock(return_value=False)
    with patch("urllib.request.urlopen", return_value=mock_resp):
        ip = get_public_ip()
    assert ip == "1.2.3.4"


def test_get_public_ip_retorna_none_en_error():
    with patch("urllib.request.urlopen", side_effect=OSError("timeout")):
        ip = get_public_ip()
    assert ip is None


@pytest.mark.parametrize("status", [403, 429])
def test_detect_block_reason_reconoce_estado_http(status):
    assert str(status) in detect_block_reason("", [status])


@pytest.mark.parametrize(
    "message",
    [
        "Error de validación ReCaptcha",
        "Demasiadas solicitudes. Intente nuevamente más tarde.",
        "Too Many Requests",
    ],
)
def test_detect_block_reason_reconoce_mensaje_del_portal(message):
    assert detect_block_reason(message) is not None


def test_detect_block_reason_ignora_respuesta_normal():
    assert detect_block_reason("Se encuentra incorporado(a) a AFP HABITAT", [200]) is None


def test_detector_http_ignora_recursos_estaticos_del_portal():
    response = MagicMock()
    response.url = "https://miportal.spensiones.gob.cl/assets/app.js"
    response.request.resource_type = "script"
    assert scraper._is_query_response(response) is False


def test_detector_http_acepta_respuesta_de_api():
    response = MagicMock()
    response.url = "https://wn1ilshgph.execute-api.us-east-1.amazonaws.com/consulta"
    response.request.resource_type = "fetch"
    assert scraper._is_query_response(response) is True


def test_query_rut_propaga_bloqueo_sin_convertirlo_en_error():
    blocked = PortalBlockedError("HTTP 429", status_code=429)
    with (
        patch.object(scraper, "_query_once", side_effect=blocked) as query_once,
        patch.object(scraper.time, "sleep"),
    ):
        with pytest.raises(PortalBlockedError, match="429"):
            scraper.query_rut(MagicMock(), "15800185-3", pause_seconds=0)

    query_once.assert_called_once()


def test_adaptive_pause_no_obliga_cambio_de_ip_en_consulta_90():
    assert worker.adaptive_pause_seconds(1.5, 89) == 1.5
    assert worker.adaptive_pause_seconds(1.5, 90) == 2.0
    assert worker.adaptive_pause_seconds(1.5, 180) == 2.5


def test_worker_reintenta_bloqueo_y_reanuda_tras_cambio_de_ip(tmp_path, monkeypatch):
    input_path = tmp_path / "entrada.xlsx"
    output_path = tmp_path / "salida.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["RUT"])
    ws.append(["15800185-3"])
    wb.save(input_path)

    emitted = []
    fake_page = MagicMock()
    fake_browser = MagicMock()
    fake_browser.new_page.return_value = fake_page
    fake_playwright = MagicMock()
    fake_playwright.chromium.launch.return_value = fake_browser
    fake_launcher = MagicMock()
    fake_launcher.start.return_value = fake_playwright

    monkeypatch.setattr(
        worker,
        "parse_args",
        lambda: Namespace(
            input=str(input_path),
            output=str(output_path),
            batch_size=1,
            pause_seconds=0,
        ),
    )
    monkeypatch.setattr(worker, "sync_playwright", lambda: fake_launcher)
    monkeypatch.setattr(worker, "get_public_ip", lambda: "1.2.3.4")
    monkeypatch.setattr(worker, "wait_for_ip_change", lambda _previous: "5.6.7.8")
    monkeypatch.setattr(worker, "BLOCK_RETRY_DELAYS", (0, 0))
    monkeypatch.setattr(worker, "STOP_REQUESTED", False)
    monkeypatch.setattr(worker, "emit", lambda event_type, **payload: emitted.append((event_type, payload)))
    query = MagicMock(
        side_effect=[
            PortalBlockedError("HTTP 429", status_code=429),
            PortalBlockedError("HTTP 429", status_code=429),
            PortalBlockedError("HTTP 429", status_code=429),
            "HABITAT",
        ]
    )
    monkeypatch.setattr(worker, "query_rut", query)

    assert worker.main() == 0
    assert query.call_count == 4
    assert any(event_type == "portal_blocked" for event_type, _ in emitted)
    assert emitted[-1][0] == "done"

    result = load_workbook(output_path)
    assert result.active["B2"].value == "HABITAT"


def test_wait_for_ip_change_reintenta_hasta_detectar_una_nueva(monkeypatch):
    commands = iter(["verify_ip\n", "verify_ip\n"])
    emitted = []

    monkeypatch.setattr(worker.sys.stdin, "readline", lambda: next(commands))
    monkeypatch.setattr(worker, "get_public_ip", MagicMock(side_effect=["1.2.3.4", "5.6.7.8"]))
    monkeypatch.setattr(worker, "emit", lambda event_type, **payload: emitted.append((event_type, payload)))
    monkeypatch.setattr(worker, "STOP_REQUESTED", False)

    assert worker.wait_for_ip_change("1.2.3.4") == "5.6.7.8"
    assert [event_type for event_type, _ in emitted] == ["ip_check", "ip_changed"]


def test_wait_for_ip_change_acepta_ip_si_no_hubo_lectura_inicial(monkeypatch):
    monkeypatch.setattr(worker.sys.stdin, "readline", lambda: "verify_ip\n")
    monkeypatch.setattr(worker, "get_public_ip", lambda: "5.6.7.8")
    monkeypatch.setattr(worker, "emit", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(worker, "STOP_REQUESTED", False)

    assert worker.wait_for_ip_change(None) == "5.6.7.8"
